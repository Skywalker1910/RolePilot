import csv
import io
import math
import re
import zipfile
from datetime import datetime
from pathlib import Path

import openpyxl
from fastapi import HTTPException
from pydantic import ValidationError

from .jobs import duplicates, job_data, write_job
from .models import ActivityEvent, ImportBatch, Job
from .schemas import JobInput, normalize

MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 2000
OFFICIAL_HEADERS = [
    "company_name",
    "job_title",
    "job_id",
    "job_url",
    "source_url",
    "source_platform",
    "date_discovered",
    "date_posted",
    "date_verified",
    "location",
    "work_mode",
    "employment_type",
    "seniority",
    "experience_min",
    "experience_max",
    "education_required",
    "required_skills",
    "preferred_skills",
    "job_keywords",
    "job_description",
    "responsibilities",
    "qualifications",
    "salary_min",
    "salary_max",
    "salary_currency",
    "salary_period",
    "sponsorship_policy",
    "citizenship_requirement",
    "security_clearance",
    "remote_eligibility",
    "application_deadline",
    "company_industry",
    "job_summary",
    "application_questions",
    "data_confidence",
    "verification_status",
    "research_notes",
]
COLUMN_FIELDS = {
    "job_id": "company_job_id",
    "job_url": "canonical_job_url",
    "work_mode": "remote_type",
    "experience_min": "minimum_experience",
    "experience_max": "maximum_experience",
    "education_required": "education_requirement",
    "job_description": "full_description",
    "date_posted": "posted_at",
    "date_verified": "source_checked_at",
    "security_clearance": "security_clearance_requirement",
}


class WorksheetSelectionRequired(ValueError):
    def __init__(self, worksheets):
        self.worksheets = worksheets
        super().__init__("This workbook has no 'jobs' worksheet. Choose a worksheet to preview.")


ALIASES = {
    "company": "company_name",
    "employer": "company_name",
    "company name": "company_name",
    "title": "job_title",
    "role": "job_title",
    "job title": "job_title",
    "url": "canonical_job_url",
    "job url": "canonical_job_url",
    "application url": "canonical_job_url",
    "direct company careers url": "canonical_job_url",
    "careers url": "canonical_job_url",
    "job link": "canonical_job_url",
    "requisition id": "company_job_id",
    "job id": "company_job_id",
    "req id": "company_job_id",
    "description": "full_description",
    "skills": "required_skills",
    "location": "location",
    "salary minimum": "salary_min",
    "salary maximum": "salary_max",
    "experience": "minimum_experience",
}
ALIASES.update({f.replace("_", " "): f for f in JobInput.model_fields})
ALIASES.update({k.replace("_", " "): v for k, v in COLUMN_FIELDS.items()})


def parse_file(
    filename: str,
    content: bytes,
    *,
    worksheet_name=None,
    include_row_numbers=False,
    include_worksheets=False,
):
    if len(content) > MAX_BYTES:
        raise ValueError("Maximum file size is 10 MB")
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        try:
            text = content.decode("utf-8-sig")
        except UnicodeError as exc:
            raise ValueError("Save CSV as UTF-8") from exc
        if "\x00" in text:
            raise ValueError("This does not look like a UTF-8 CSV file")
        rows = list(csv.reader(io.StringIO(text)))
        worksheets = []
        selected_sheet = None
    elif suffix == ".xlsx":
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                names = set(archive.namelist())
                if "[Content_Types].xml" not in names or "xl/workbook.xml" not in names:
                    raise ValueError("This ZIP file is not a valid .xlsx workbook")
                if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
                    raise ValueError("Workbook expands beyond the safe 100 MB limit")
        except zipfile.BadZipFile as exc:
            raise ValueError("The .xlsx file is not a valid Excel workbook") from exc
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=False)
        try:
            worksheets = list(workbook.sheetnames)
            if worksheet_name is None:
                if "jobs" in worksheets:
                    selected_sheet = "jobs"
                else:
                    raise WorksheetSelectionRequired(worksheets)
            elif worksheet_name not in worksheets:
                raise ValueError("Selected worksheet is not present in the workbook")
            else:
                selected_sheet = worksheet_name
            sheet = workbook[selected_sheet]
            rows = []
            for row in sheet.iter_rows():
                if len(rows) > MAX_ROWS:
                    raise ValueError(f"Maximum {MAX_ROWS} rows per import")
                if len(row) > 100:
                    raise ValueError("Maximum 100 columns")
                rows.append([cell.value for cell in row])
        finally:
            workbook.close()
    else:
        raise ValueError("Choose an .xlsx or .csv file")
    if not rows:
        raise ValueError("The file is empty")
    headers = [str(v or "").strip() for v in rows[0]]
    while headers and not headers[-1]:
        headers.pop()
    if not headers or len(headers) > 100 or any(not h for h in headers):
        raise ValueError("Every column needs a header (maximum 100 columns)")
    if len(set(normalize(h) for h in headers)) != len(headers):
        raise ValueError("Duplicate column headers are not supported")

    def cell(v):
        if isinstance(v, datetime):
            return v.isoformat()
        return str(v) if v is not None else ""

    result = []
    row_numbers = []
    for source_row_number, row in enumerate(rows[1:], start=2):
        if any(v is not None and str(v).strip() for v in row):
            if len(row) > len(headers) and any(str(v or "").strip() for v in row[len(headers) :]):
                raise ValueError("A row has values beyond its headers")
            converted = {}
            for i, header in enumerate(headers):
                raw_value = row[i] if i < len(row) else None
                value = cell(raw_value)
                field = ALIASES.get(normalize(header).replace("_", " "))
                if field == "company_job_id" and isinstance(raw_value, (int, float)):
                    numeric_id = raw_value
                    if isinstance(numeric_id, float) and numeric_id.is_integer():
                        value = str(int(numeric_id))
                converted[header] = value
            result.append(converted)
            row_numbers.append(source_row_number)
    if not result or len(result) > MAX_ROWS:
        raise ValueError(f"Provide between 1 and {MAX_ROWS} non-empty rows")
    mapping = {h: ALIASES.get(normalize(h).replace("_", " "), "") for h in headers}
    if include_worksheets:
        return result, mapping, row_numbers, worksheets, selected_sheet
    if include_row_numbers:
        return result, mapping, row_numbers
    return result, mapping


def number(value):
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        if not math.isfinite(value):
            raise ValueError("Use a finite number")
        return value
    cleaned = re.sub(r"[$,€£]", "", value.strip().lower())
    if re.fullmatch(r"\d+(\.\d+)?k", cleaned):
        return float(cleaned[:-1]) * 1000
    parsed = float(cleaned)
    if not math.isfinite(parsed):
        raise ValueError("Use a finite number")
    return parsed


def review(session, preview, mapping, corrections):
    fields = [v for v in mapping.values() if v]
    if len(fields) != len(set(fields)):
        raise HTTPException(422, "Map each field only once")
    if set(fields) - set(JobInput.model_fields):
        raise HTTPException(422, "Unsupported mapped field")
    if set(mapping) - set(preview.rows[0]):
        raise HTTPException(422, "Unknown source column")
    results = []
    seen = []
    row_numbers = preview.row_numbers or list(range(2, len(preview.rows) + 2))
    for index, raw in zip(row_numbers, preview.rows, strict=True):
        values = {field: raw.get(column, "") for column, field in mapping.items() if field}
        values.update(corrections.get(index, {}))
        errors, warnings = [], []
        data = None
        try:
            for field in ("salary_min", "salary_max", "minimum_experience", "maximum_experience"):
                if field in values:
                    values[field] = number(values[field])
            data = JobInput.model_validate(values)
        except ValidationError as exc:
            errors = [
                f"{'.'.join(str(v) for v in e['loc']) or 'row'}: {e['msg']}" for e in exc.errors()
            ]
        except (ValueError, TypeError):
            errors = [
                "Salary and experience must be numbers; split ranges into minimum and maximum"
            ]
        candidates, kind, changes = [], None, {}
        if data:
            kind, jobs = duplicates(session, data)
            candidates = [
                {
                    "id": j.id,
                    "revision": j.revision,
                    "company_name": j.company.name,
                    "job_title": j.job_title,
                }
                for j in jobs
            ]
            for job in jobs:
                before = job_data(job)
                changes[str(job.id)] = {
                    key: {
                        "current": before[key],
                        "incoming": value,
                        "locally_edited": key in job.locally_edited_fields,
                    }
                    for key, value in data.model_dump().items()
                    if before[key] != value
                }
            for previous_number, previous in seen:
                exact = data.canonical_job_url == previous.canonical_job_url or (
                    data.company_job_id
                    and previous.company_job_id
                    and normalize(data.company_name) == normalize(previous.company_name)
                    and normalize(data.company_job_id) == normalize(previous.company_job_id)
                )
                possible = all(
                    normalize(getattr(data, f) or "") == normalize(getattr(previous, f) or "")
                    for f in ("company_name", "job_title", "city", "state", "country")
                )
                if exact:
                    errors.append(
                        f"Duplicate identity in this file: row {previous_number}; skip one"
                    )
                elif possible:
                    warnings.append(
                        f"Possible duplicate of row {previous_number}; verify separately"
                    )
            seen.append((index, data))
            host = data.canonical_job_url.split("/")[2]
            if any(h in host for h in ("linkedin.com", "indeed.com", "glassdoor.com", "bit.ly")):
                warnings.append(
                    "This may be an aggregator or shortened URL. Verify the employer link."
                )
            warnings.append("Employer ownership and listing accuracy must be verified by you.")
        results.append(
            {
                "row_number": index,
                "raw": raw,
                "data": data.model_dump() if data else values,
                "errors": errors,
                "warnings": warnings,
                "duplicate_kind": kind,
                "candidates": candidates,
                "conflicts": changes,
            }
        )
    preview.mapping = mapping
    preview.reviewed_rows = results
    return results


def commit(session, preview, decisions):
    if preview.committed_batch_id:
        raise HTTPException(409, "This preview has already been imported")
    reviewed = {r["row_number"]: r for r in preview.reviewed_rows}
    if (
        not reviewed
        or len(decisions) != len(reviewed)
        or {d.row_number for d in decisions} != set(reviewed)
    ):
        raise HTTPException(422, "Choose one action for every reviewed row")
    batch = ImportBatch(
        filename=preview.filename,
        rows_processed=len(reviewed),
        inserted=0,
        updated=0,
        skipped=0,
        rejected=0,
        source_metadata={"mapping": preview.mapping},
        validation_errors=[],
    )
    session.add(batch)
    session.flush()
    validation_errors = []
    for decision in decisions:
        row = reviewed[decision.row_number]
        if decision.action == "skip":
            batch.skipped += 1
            if row["errors"]:
                batch.rejected += 1
                validation_errors.append({"row": decision.row_number, "errors": row["errors"]})
            continue
        if row["errors"]:
            raise HTTPException(422, f"Row {decision.row_number} has validation errors")
        data = JobInput.model_validate(row["data"])
        kind, matches = duplicates(session, data)
        job = None
        if decision.action == "insert":
            if kind == "exact":
                raise HTTPException(409, "An exact duplicate now exists; review again")
            if kind == "possible" and not decision.approve_conflicts:
                raise HTTPException(409, "Confirm this possible duplicate is a separate job")
            batch.inserted += 1
        else:
            candidate_ids = {c["id"] for c in row["candidates"]}
            if decision.target_id not in candidate_ids:
                raise HTTPException(409, "Update target was not part of the reviewed candidates")
            reviewed_revision = next(
                c["revision"] for c in row["candidates"] if c["id"] == decision.target_id
            )
            if decision.expected_revision != reviewed_revision:
                raise HTTPException(409, "Revision does not match the reviewed candidate")
            job = session.get(Job, decision.target_id)
            if not job or job.revision != decision.expected_revision:
                raise HTTPException(409, "Job changed since preview; review again")
            if kind == "exact" and any(j.id != job.id for j in matches):
                raise HTTPException(
                    409, "Conflicting requisition and URL identities; correct the row"
                )
            if not decision.approve_conflicts:
                raise HTTPException(409, "Explicitly approve the displayed field replacements")
            batch.updated += 1
        write_job(
            session, data, job, batch_id=batch.id, raw=row["raw"], row_number=decision.row_number
        )
    batch.validation_errors = validation_errors
    preview.committed_batch_id = batch.id
    session.add(
        ActivityEvent(
            event_type="import.completed",
            message=f"{batch.filename}: {batch.inserted} inserted, "
            f"{batch.updated} updated, {batch.skipped} skipped",
        )
    )
    session.flush()
    return batch


def batch_data(batch):
    return {c.name: getattr(batch, c.name) for c in ImportBatch.__table__.columns}
