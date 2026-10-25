import hmac
import io
import os
import uuid
from contextlib import asynccontextmanager
from pathlib import Path
from typing import Annotated

import openpyxl
from fastapi import Depends, FastAPI, Form, HTTPException, Query, Request, UploadFile
from fastapi.responses import JSONResponse, Response
from openpyxl.styles import Font, PatternFill
from sqlalchemy import func, or_, select, text
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session, selectinload
from starlette.concurrency import run_in_threadpool
from starlette.middleware.trustedhost import TrustedHostMiddleware

from . import db
from .admin import router as admin_router
from .applications_api import application_metrics
from .applications_api import router as applications_router
from .auth import router as auth_router
from .auth import session_from_request, session_user
from .auth_models import AuthUser
from .db import get_session, utcnow
from .imports import MAX_BYTES, WorksheetSelectionRequired, batch_data, commit, parse_file, review
from .jobs import duplicates, serialize, write_job
from .library import ensure_library, scan_library
from .library_api import router as library_router
from .models import ActivityEvent, Company, ImportBatch, ImportPreview, Job, JobSkill, Skill
from .production_config import validate_production_environment
from .schemas import (
    ArchiveRequest,
    CommitRequest,
    CompanyCreate,
    CompanyPatch,
    EditRequest,
    JobInput,
    ListingStatusRequest,
    ReviewRequest,
    ShortlistRequest,
)

DB = Annotated[Session, Depends(get_session)]


@asynccontextmanager
async def lifespan(app):
    # Startup scanning is read-only for source files, transactional for the index.
    def startup_scan():
        with db.SessionLocal() as session:
            session.execute(text("BEGIN IMMEDIATE"))
            if app.state.auth_required:
                if (
                    session.scalar(
                        select(AuthUser.id)
                        .where(AuthUser.role == "admin", AuthUser.is_active.is_(True))
                        .limit(1)
                    )
                    is None
                ):
                    raise RuntimeError(
                        "Provision the first administrator with `python -m app.auth_cli init-user`"
                    )
            else:
                scan_library(session, ensure_library(session))
            session.commit()

    await run_in_threadpool(startup_scan)
    yield


app = FastAPI(title="RolePilot local API", version="0.2.0", lifespan=lifespan)
app.state.auth_required = os.environ.get("ROLEPILOT_AUTH_REQUIRED", "0").casefold() in {
    "1",
    "true",
    "yes",
}
app.state.secure_cookies = os.environ.get(
    "ROLEPILOT_COOKIE_SECURE", "1" if app.state.auth_required else "0"
).casefold() in {"1", "true", "yes"}
app.state.public_origin = os.environ.get("ROLEPILOT_PUBLIC_ORIGIN", "https://career.adityamore.dev")
if app.state.auth_required:
    try:
        validate_production_environment()
    except ValueError as exc:
        raise RuntimeError(str(exc)) from exc
if app.state.auth_required and not app.state.public_origin.startswith("https://"):
    raise RuntimeError("ROLEPILOT_PUBLIC_ORIGIN must use HTTPS when authentication is enabled")
if app.state.auth_required and not app.state.secure_cookies:
    raise RuntimeError("Secure authentication cookies are required when authentication is enabled")
allowed_hosts = {"127.0.0.1", "localhost", "testserver"}
allowed_hosts.update(
    host.strip()
    for host in os.environ.get("ROLEPILOT_ALLOWED_HOSTS", "career.adityamore.dev").split(",")
    if host.strip()
)
app.include_router(library_router)
app.include_router(applications_router)
app.include_router(auth_router)
app.include_router(admin_router)
app.add_middleware(TrustedHostMiddleware, allowed_hosts=sorted(allowed_hosts))


@app.middleware("http")
async def local_requests(request: Request, call_next):
    path = request.url.path
    if path.startswith("/api") and app.state.auth_required:
        record = None
        with db.SessionLocal() as session:
            record = session_from_request(request, session)
            user = session_user(record, session)
        request.state.auth_session = record
        request.state.auth_user = user
        public_paths = {
            "/api/health",
            "/api/auth/login",
            "/api/auth/signup",
            "/api/auth/session",
        }
        if path not in public_paths and record is None:
            return JSONResponse({"detail": "Authentication required"}, 401)
        if record is not None and request.state.auth_user.role == "admin":
            admin_paths = {
                "/api/auth/session",
                "/api/auth/logout",
                "/api/auth/profile",
                "/api/health",
            }
            if path not in admin_paths and not path.startswith("/api/admin/"):
                return JSONResponse(
                    {"detail": "Administrator accounts cannot access job workspace data"}, 403
                )
        if request.method not in {"GET", "HEAD", "OPTIONS"}:
            if request.headers.get("origin") != app.state.public_origin:
                return JSONResponse({"detail": "Request origin is not allowed"}, 403)
            if path not in {"/api/auth/login", "/api/auth/signup"}:
                csrf = request.headers.get("x-csrf-token", "")
                if not record or not hmac.compare_digest(csrf, record.csrf_token):
                    return JSONResponse({"detail": "CSRF token is missing or invalid"}, 403)
    elif request.method not in {"GET", "HEAD", "OPTIONS"}:
        origin = request.headers.get("origin")
        if origin and origin not in {
            "http://127.0.0.1:5173",
            "http://127.0.0.1:5174",
            "http://localhost:5173",
            "http://localhost:5174",
            "http://127.0.0.1:8000",
        }:
            return JSONResponse({"detail": "Only the local RolePilot UI may write data"}, 403)
        if request.headers.get("x-rolepilot") != "local":
            return JSONResponse({"detail": "Missing local request header"}, 403)
    response = await call_next(request)
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["Referrer-Policy"] = "no-referrer"
    return response


@app.exception_handler(IntegrityError)
async def integrity_error(request, exc):
    return JSONResponse(
        {"detail": "Identity conflict or invalid relation; refresh and review again"}, 409
    )


@app.get("/api/health")
def health(session: DB):
    version = session.execute(text("SELECT version_num FROM alembic_version")).scalar_one()
    return {"status": "ok", "schema_version": version}


@app.get("/api/fields")
def fields():
    return list(JobInput.model_fields)


@app.get("/api/jobs")
def jobs(
    session: DB,
    q: str = "",
    company: str = "",
    location: str = "",
    skill: str = "",
    remote_type: str = "",
    sponsorship: str = "",
    salary_min: float | None = None,
    experience_max: float | None = None,
    company_id: int | None = None,
    employment_type: str = "",
    seniority: str = "",
    experience_min: float | None = None,
    include_unspecified_experience: bool = False,
    salary_max: float | None = None,
    preferred_skill: str = "",
    citizenship_requirement: str = "",
    security_clearance: str = "",
    verification_status: str = "",
    date_posted_from: str = "",
    date_posted_to: str = "",
    date_imported_from: str = "",
    date_imported_to: str = "",
    shortlisted: bool | None = None,
    archived: bool = False,
    sort: str = Query("newest", pattern="^(newest|title|salary|oldest|company|posted)$"),
    page: int = Query(1, ge=1),
    page_size: int = Query(25, ge=1, le=100),
):
    stmt = select(Job).join(Company)
    stmt = stmt.where(Job.archived_at.is_not(None) if archived else Job.archived_at.is_(None))
    if q:
        stmt = stmt.where(
            or_(
                Job.job_title.contains(q, autoescape=True),
                Company.name.contains(q, autoescape=True),
                Job.full_description.contains(q, autoescape=True),
                Job.company_job_id.contains(q, autoescape=True),
                Job.skills.any(
                    JobSkill.skill.has(
                        Skill.normalized_name.contains(q.casefold(), autoescape=True)
                    )
                ),
            )
        )
    if company:
        stmt = stmt.where(Company.name.contains(company, autoescape=True))
    if company_id is not None:
        stmt = stmt.where(Job.company_id == company_id)
    if location:
        stmt = stmt.where(
            or_(
                *[
                    getattr(Job, f).contains(location, autoescape=True)
                    for f in ("city", "state", "country", "location")
                ]
            )
        )
    if skill:
        stmt = stmt.where(
            Job.skills.any(
                JobSkill.skill.has(
                    Skill.normalized_name.contains(skill.casefold(), autoescape=True)
                )
            )
        )
    if preferred_skill:
        stmt = stmt.where(
            Job.skills.any(
                (JobSkill.kind == "Preferred")
                & JobSkill.skill.has(
                    Skill.normalized_name.contains(preferred_skill.casefold(), autoescape=True)
                )
            )
        )
    if remote_type:
        stmt = stmt.where(Job.remote_type == remote_type)
    if sponsorship:
        stmt = stmt.where(Job.sponsorship_policy.contains(sponsorship, autoescape=True))
    if salary_min is not None:
        stmt = stmt.where(Job.salary_max >= salary_min)
    if experience_max is not None:
        clause = Job.minimum_experience <= experience_max
        stmt = stmt.where(
            or_(clause, Job.minimum_experience.is_(None))
            if include_unspecified_experience
            else clause
        )
    if experience_min is not None:
        stmt = stmt.where(Job.maximum_experience >= experience_min)
    if salary_max is not None:
        stmt = stmt.where(Job.salary_min <= salary_max)
    if employment_type:
        stmt = stmt.where(Job.employment_type == employment_type)
    if seniority:
        stmt = stmt.where(Job.seniority == seniority)
    if citizenship_requirement:
        stmt = stmt.where(
            Job.citizenship_requirement.contains(citizenship_requirement, autoescape=True)
        )
    if security_clearance:
        stmt = stmt.where(
            Job.security_clearance_requirement.contains(security_clearance, autoescape=True)
        )
    if verification_status:
        stmt = stmt.where(Job.verification_status == verification_status)
    if date_posted_from:
        stmt = stmt.where(Job.posted_at >= date_posted_from)
    if date_posted_to:
        stmt = stmt.where(Job.posted_at <= date_posted_to)
    if date_imported_from:
        stmt = stmt.where(Job.imported_at >= date_imported_from)
    if date_imported_to:
        stmt = stmt.where(Job.imported_at <= date_imported_to + "T23:59:59")
    if shortlisted is not None:
        stmt = stmt.where(Job.is_shortlisted.is_(shortlisted))
    total = session.scalar(select(func.count()).select_from(stmt.subquery()))
    order = {
        "newest": Job.imported_at.desc(),
        "title": Job.job_title.asc(),
        "salary": Job.salary_max.desc(),
        "oldest": Job.imported_at.asc(),
        "company": Company.name.asc(),
        "posted": Job.posted_at.desc(),
    }[sort]
    rows = session.scalars(
        stmt.options(
            selectinload(Job.company), selectinload(Job.skills).selectinload(JobSkill.skill)
        )
        .order_by(order, Job.id.desc())
        .offset((page - 1) * page_size)
        .limit(page_size)
    )
    return {
        "items": [serialize(j) for j in rows],
        "total": total,
        "page": page,
        "page_size": page_size,
    }


def get_job(session, job_id):
    job = session.get(Job, job_id)
    if not job:
        raise HTTPException(404, "Job not found")
    return job


@app.get("/api/jobs/{job_id}")
def detail(job_id: int, session: DB):
    return serialize(get_job(session, job_id), detail=True)


@app.post("/api/jobs", status_code=201)
def create_job(body: JobInput, session: DB):
    kind, matches = duplicates(session, body)
    if matches:
        raise HTTPException(409, "A matching job exists; use import review to resolve duplicates")
    job = write_job(session, body, local=True)
    session.commit()
    return serialize(job, detail=True)


@app.put("/api/jobs/{job_id}")
def edit_job(job_id: int, body: EditRequest, session: DB):
    # Serialize writes before reading revisions, including separate processes.
    session.execute(text("BEGIN IMMEDIATE"))
    job = get_job(session, job_id)
    if job.revision != body.expected_revision:
        raise HTTPException(409, "Job changed; reopen it before saving")
    kind, matches = duplicates(session, body.job)
    if kind == "exact" and any(j.id != job.id for j in matches):
        raise HTTPException(409, "Another job has that URL or requisition ID")
    write_job(session, body.job, job, local=True)
    session.commit()
    return serialize(job, detail=True)


@app.patch("/api/jobs/{job_id}")
def patch_job(job_id: int, body: EditRequest, session: DB):
    return edit_job(job_id, body, session)


@app.delete("/api/jobs/{job_id}")
def delete_job(job_id: int, expected_revision: int, session: DB):
    # Removal is deliberately reversible: preserve application/source history in Archive.
    session.execute(text("BEGIN IMMEDIATE"))
    job = get_job(session, job_id)
    if job.revision != expected_revision:
        raise HTTPException(409, "Job changed; refresh before removing it")
    if job.archived_at is None or job.listing_status != "Removed":
        job.archived_at = job.archived_at or utcnow()
        job.listing_status = "Removed"
        job.revision += 1
        session.add(
            ActivityEvent(
                event_type="job.removed",
                message=f"{job.company.name} · {job.job_title}",
            )
        )
        session.commit()
    return serialize(job)


@app.patch("/api/jobs/{job_id}/listing-status")
def update_listing_status(job_id: int, body: ListingStatusRequest, session: DB):
    session.execute(text("BEGIN IMMEDIATE"))
    job = get_job(session, job_id)
    if job.revision != body.expected_revision:
        raise HTTPException(409, "Job changed; refresh before updating its listing status")
    job.listing_status = body.listing_status
    if body.listing_status == "Unavailable":
        job.archived_at = utcnow()
    elif body.listing_status == "Available":
        job.archived_at = None
    elif body.listing_status == "Removed":
        job.archived_at = job.archived_at or utcnow()
    job.revision += 1
    session.add(
        ActivityEvent(
            event_type=f"job.listing_status.{body.listing_status.casefold()}",
            message=f"{job.company.name} · {job.job_title}",
        )
    )
    session.commit()
    return serialize(job)


@app.post("/api/jobs/{job_id}/archive")
def archive_job(job_id: int, body: ArchiveRequest, session: DB):
    session.execute(text("BEGIN IMMEDIATE"))
    job = get_job(session, job_id)
    if job.revision != body.expected_revision:
        raise HTTPException(409, "Job changed; refresh before archiving")
    job.archived_at = utcnow() if body.archived else None
    if not body.archived and job.listing_status in {"Unavailable", "Removed"}:
        job.listing_status = "Available"
    job.revision += 1
    session.add(
        ActivityEvent(
            event_type="job.archived" if body.archived else "job.restored",
            message=f"{job.company.name} · {job.job_title}",
        )
    )
    session.commit()
    return serialize(job)


@app.patch("/api/jobs/{job_id}/shortlist")
def shortlist_job(job_id: int, body: ShortlistRequest, session: DB):
    session.execute(text("BEGIN IMMEDIATE"))
    job = get_job(session, job_id)
    if job.revision != body.expected_revision:
        raise HTTPException(409, "Job changed; refresh before shortlisting")
    job.is_shortlisted = body.shortlisted
    job.revision += 1
    session.add(
        ActivityEvent(
            event_type="job.shortlist.updated", message=f"{job.company.name} · {job.job_title}"
        )
    )
    session.commit()
    return serialize(job)


@app.get("/api/companies")
def companies(session: DB):
    active_job_counts = dict(
        session.execute(
            select(Job.company_id, func.count(Job.id))
            .where(Job.archived_at.is_(None))
            .group_by(Job.company_id)
        ).all()
    )
    return [
        {
            "id": c.id,
            "name": c.name,
            "industry": c.industry,
            "website": c.website,
            "careers_page_url": c.careers_page_url,
            "headquarters": c.headquarters,
            "notes": c.notes,
            "created_at": c.created_at,
            "updated_at": c.updated_at,
            "active_job_count": active_job_counts.get(c.id, 0),
        }
        for c in session.scalars(select(Company).order_by(Company.name))
    ]


@app.get("/api/companies/{company_id}")
def company_detail(company_id: int, session: DB):
    company = session.get(Company, company_id)
    if not company:
        raise HTTPException(404, "Company not found")
    return {
        "id": company.id,
        "name": company.name,
        "normalized_name": company.normalized_name,
        "industry": company.industry,
        "website": company.website,
        "careers_page_url": company.careers_page_url,
        "headquarters": company.headquarters,
        "notes": company.notes,
        "created_at": company.created_at,
        "updated_at": company.updated_at,
        "jobs": [serialize(j) for j in company.jobs],
    }


@app.post("/api/companies", status_code=201)
def create_company(body: CompanyCreate, session: DB):
    from .schemas import normalize

    key = normalize(body.name)
    if session.scalar(select(Company).where(Company.normalized_name == key)):
        raise HTTPException(409, "A company with this normalized name already exists")
    company = Company(**body.model_dump(), normalized_name=key)
    session.add(company)
    session.commit()
    return {
        "id": company.id,
        "name": company.name,
        "industry": company.industry,
        "website": company.website,
        "careers_page_url": company.careers_page_url,
        "headquarters": company.headquarters,
        "notes": company.notes,
        "created_at": company.created_at,
        "updated_at": company.updated_at,
    }


@app.patch("/api/companies/{company_id}")
def patch_company(company_id: int, body: CompanyPatch, session: DB):
    from .schemas import normalize

    session.execute(text("BEGIN IMMEDIATE"))
    company = session.get(Company, company_id)
    if not company:
        raise HTTPException(404, "Company not found")
    if company.updated_at != body.expected_revision:
        raise HTTPException(409, "Company changed; reload before saving")
    values = body.model_dump(exclude={"expected_revision"}, exclude_unset=True)
    if "name" in values:
        company.normalized_name = normalize(values["name"])
    for key, value in values.items():
        setattr(company, key, value)
    session.commit()
    return {
        "id": company.id,
        "name": company.name,
        "industry": company.industry,
        "website": company.website,
        "careers_page_url": company.careers_page_url,
        "headquarters": company.headquarters,
        "notes": company.notes,
        "created_at": company.created_at,
        "updated_at": company.updated_at,
    }


@app.get("/api/skills")
def skills(session: DB, q: str = "", limit: int = Query(100, ge=1, le=500)):
    stmt = select(Skill)
    if q:
        stmt = stmt.where(Skill.normalized_name.contains(q.casefold(), autoescape=True))
    return [
        {
            "id": s.id,
            "name": s.canonical_name or s.name,
            "normalized_name": s.normalized_name,
            "jobs": session.scalar(
                select(func.count()).select_from(JobSkill).where(JobSkill.skill_id == s.id)
            ),
        }
        for s in session.scalars(stmt.order_by(Skill.normalized_name).limit(limit))
    ]


@app.get("/api/overview")
def overview(session: DB):
    active = Job.archived_at.is_(None)
    recent = session.scalars(select(Job).where(active).order_by(Job.id.desc()).limit(5))
    company_counts = session.execute(
        select(Company.name, func.count(Job.id))
        .join(Job)
        .where(active)
        .group_by(Company.id)
        .order_by(func.count(Job.id).desc())
        .limit(6)
    )
    return {
        "total_jobs": session.scalar(select(func.count(Job.id)).where(active)),
        "archived_jobs": session.scalar(
            select(func.count(Job.id)).where(Job.archived_at.is_not(None))
        ),
        "companies": session.scalar(
            select(func.count(func.distinct(Job.company_id))).where(active)
        ),
        "import_batches": session.scalar(select(func.count(ImportBatch.id))),
        "recent_jobs": [serialize(j) for j in recent],
        "companies_chart": [{"name": name, "jobs": count} for name, count in company_counts],
        "activity": [
            {"id": e.id, "message": e.message, "created_at": e.created_at}
            for e in session.scalars(
                select(ActivityEvent).order_by(ActivityEvent.id.desc()).limit(8)
            )
        ],
        "application_metrics": application_metrics(session),
    }


@app.post("/api/imports/preview")
async def upload(file: UploadFile, session: DB, worksheet_name: str | None = Form(default=None)):
    content = await file.read(MAX_BYTES + 1)
    try:
        rows, mapping, row_numbers, worksheets, selected_sheet = parse_file(
            file.filename or "", content, worksheet_name=worksheet_name, include_worksheets=True
        )
    except WorksheetSelectionRequired as exc:
        raise HTTPException(
            409, {"message": str(exc), "worksheets": exc.worksheets, "requires_selection": True}
        ) from exc
    except Exception as exc:
        if isinstance(exc, HTTPException):
            raise
        raise HTTPException(
            422,
            str(exc)
            if isinstance(exc, ValueError)
            else "Cannot read this workbook; save it as a new .xlsx file",
        ) from exc
    preview = ImportPreview(
        id=str(uuid.uuid4()),
        filename=Path(file.filename or "import").name,
        rows=rows,
        row_numbers=row_numbers,
        mapping=mapping,
    )
    session.add(preview)
    # Duplicate auto-mappings are left for the user to resolve.
    used = set()
    for column, field in mapping.items():
        if field in used:
            mapping[column] = ""
        elif field:
            used.add(field)
    result = review(session, preview, mapping, {})
    session.commit()
    return {
        "id": preview.id,
        "filename": preview.filename,
        "worksheets": worksheets,
        "worksheet_name": selected_sheet,
        "mapping": mapping,
        "rows": result,
        "fields": list(JobInput.model_fields),
    }


def get_preview(session, preview_id):
    preview = session.get(ImportPreview, preview_id)
    if not preview:
        raise HTTPException(404, "Preview not found")
    if preview.committed_batch_id:
        raise HTTPException(409, "Preview already committed")
    return preview


@app.post("/api/imports/{preview_id}/review")
def review_import(preview_id: str, body: ReviewRequest, session: DB):
    preview = get_preview(session, preview_id)
    rows = review(session, preview, body.mapping, body.corrections)
    session.commit()
    return {
        "id": preview.id,
        "filename": preview.filename,
        "mapping": preview.mapping,
        "rows": rows,
        "fields": list(JobInput.model_fields),
    }


@app.post("/api/imports/{preview_id}/commit")
def commit_import(preview_id: str, body: CommitRequest, session: DB):
    session.execute(text("BEGIN IMMEDIATE"))
    preview = get_preview(session, preview_id)
    batch = commit(session, preview, body.decisions)
    session.commit()
    return batch_data(batch)


@app.get("/api/imports")
def batches(session: DB):
    return [
        batch_data(b)
        for b in session.scalars(select(ImportBatch).order_by(ImportBatch.id.desc()).limit(100))
    ]


def spreadsheet_safe(value):
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + value
    return value


@app.get("/api/template.xlsx")
def template(sample: bool = False):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "jobs"
    columns = [
        "company_name",
        "job_title",
        "job_id",
        "job_url",
        "source_url",
        "source_platform",
        "date_discovered",
        "date_posted",
        "date_verified",
        "location",
        "work_mode",
        "employment_type",
        "seniority",
        "experience_min",
        "experience_max",
        "education_required",
        "required_skills",
        "preferred_skills",
        "job_keywords",
        "job_description",
        "responsibilities",
        "qualifications",
        "salary_min",
        "salary_max",
        "salary_currency",
        "salary_period",
        "sponsorship_policy",
        "citizenship_requirement",
        "security_clearance",
        "remote_eligibility",
        "application_deadline",
        "company_industry",
        "job_summary",
        "application_questions",
        "data_confidence",
        "verification_status",
        "research_notes",
    ]
    sheet.append(columns)
    if sample:
        example = {
            "company_name": "Juniper Orbit (fictional)",
            "job_title": "Backend Engineer",
            "job_id": "ENG-0042",
            "job_url": "https://juniper-orbit.example/jobs/eng-42",
            "location": "Raleigh, NC, US",
            "work_mode": "Hybrid",
            "employment_type": "Full-time",
            "experience_min": 2,
            "experience_max": 5,
            "required_skills": "Python; SQL; FastAPI",
            "preferred_skills": "React; Docker",
            "job_description": "Fictional listing for testing. Not a real vacancy.",
            "salary_min": 100000,
            "salary_max": 140000,
            "salary_currency": "USD",
            "salary_period": "year",
            "date_posted": "2026-10-01",
            "data_confidence": "High",
            "verification_status": "Unverified",
        }
        sheet.append([spreadsheet_safe(example.get(column)) for column in columns])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="155E75")
        sheet.column_dimensions[cell.column_letter].width = 26
    for letter, width in {"D": 58, "T": 72, "Q": 36, "R": 36, "AK": 40}.items():
        sheet.column_dimensions[letter].width = width
    for row in range(2, 1002):
        sheet.cell(row, 3).number_format = "@"
        for col in (7, 8, 9, 31):
            sheet.cell(row, col).number_format = "yyyy-mm-dd"
    info = workbook.create_sheet("metadata")
    info.append(["schema_version", "1.0"])
    info.append(["required_fields", "company_name; job_title; job_url"])
    info.append(["instructions", "Leave unknown values blank; separate skills with semicolons."])
    info.column_dimensions["A"].width = 24
    info.column_dimensions["B"].width = 90
    output = io.BytesIO()
    workbook.save(output)
    return Response(
        output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": 'attachment; filename="rolepilot-'
            + ("sample" if sample else "template")
            + '.xlsx"'
        },
    )
